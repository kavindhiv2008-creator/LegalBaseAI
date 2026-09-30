from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from backend.utils.text import (
    sanitize_text,
    split_sections,
    safe_filename,
)


ROOT = Path(__file__).resolve().parents[2]

LOGO_PATH = (
    ROOT /
    "assets" /
    "logo.png"
)


def export_txt(
    content: str
) -> bytes:

    return sanitize_text(
        content
    ).encode("utf-8")


def _set_doc_defaults(
    doc: Document
):

    style = doc.styles["Normal"]

    style.font.name = (
        "Times New Roman"
    )

    style.font.size = Pt(11)

    for section in doc.sections:

        section.top_margin = Inches(
            0.75
        )

        section.bottom_margin = Inches(
            0.75
        )

        section.left_margin = Inches(
            0.85
        )

        section.right_margin = Inches(
            0.85
        )


def export_docx(
    document_type: str,
    content: str
) -> bytes:

    doc = Document()

    _set_doc_defaults(doc)

    # Logo
    if LOGO_PATH.exists():

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run().add_picture(
            str(LOGO_PATH),
            width=Inches(2.1)
        )

    # Title
    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        sanitize_text(
            document_type
        ).upper()
    )

    run.bold = True

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(15)

    # Body
    sections = split_sections(
        content
    )

    for heading, body in sections:

        paragraph = doc.add_paragraph()

        heading_run = paragraph.add_run(
            heading
        )

        heading_run.bold = True

        heading_run.font.name = (
            "Times New Roman"
        )

        heading_run.font.size = Pt(12)

        for line in body.split("\n"):

            paragraph = doc.add_paragraph(
                line
            )

            paragraph.paragraph_format.space_after = Pt(
                6
            )

            paragraph.paragraph_format.line_spacing = (
                1.15
            )

    # Footer
    for section in doc.sections:

        footer = (
            section
            .footer
            .paragraphs[0]
        )

        footer.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        footer_run = footer.add_run(
            "LegalEase | "
            "AI-generated legal document draft"
        )

        footer_run.font.size = Pt(8)

    buffer = BytesIO()

    doc.save(buffer)

    return buffer.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        document_type: str
    ):

        super().__init__()

        self.document_type = (
            sanitize_text(
                document_type
            )
        )

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=85,
                    y=8,
                    w=40
                )

                self.set_y(28)

            except Exception:

                self.set_y(12)

        else:

            self.set_y(12)

        self.set_font(
            "Helvetica",
            "B",
            13
        )

        self.cell(
            0,
            8,
            self.document_type.upper(),
            align="C"
        )

        self.ln(10)

    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            6,
            "LegalEase | "
            "AI-generated legal document draft",
            align="C"
        )


def export_pdf(
    document_type: str,
    content: str
) -> bytes:

    pdf = LegalEasePDF(
        document_type
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "",
        10.5
    )

    sections = split_sections(
        content
    )

    for heading, body in sections:

        pdf.set_font(
            "Helvetica",
            "B",
            11
        )

        pdf.multi_cell(
            0,
            7,
            sanitize_text(
                heading
            )
        )

        pdf.ln(1)

        pdf.set_font(
            "Helvetica",
            "",
            10.5
        )

        for paragraph in body.split("\n"):

            if paragraph.strip():

                pdf.multi_cell(
                    0,
                    6,
                    sanitize_text(
                        paragraph
                    )
                )

                pdf.ln(1)

        pdf.ln(2)

    raw = pdf.output()

    return bytes(raw)


def build_filename(
    document_type: str,
    extension: str
) -> str:

    return (
        f"{safe_filename(document_type)}"
        f".{extension}"
    )